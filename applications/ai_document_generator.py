"""
JOB HUNTER BELGIUM
AI DOCUMENT GENERATOR - VERSION 1.0

Pipeline :
    latest job_refresh_v1_*.json
        + latest application_recheck_v1_*.json
        + CV local
        -> OpenAI Responses API
        -> Structured Output JSON
        -> CV DOCX + lettre DOCX + email TXT localement

IMPORTANT :
- Ne touche pas Matcher / Gate / Queue / Canonical / DB.
- Ne traite par défaut QUE READY_DOCUMENTS.
- La clé API vient uniquement de OPENAI_API_KEY.
- Aucune clé n'est écrite dans les logs.
- Les données professionnelles envoyées à l'API sont limitées au CV/profil
  et à la description de l'offre nécessaire à la candidature.

Usage :
    python -m applications.ai_document_generator --limit 1
    python -m applications.ai_document_generator

Optionnel :
    --model gpt-5.6-terra
    --recheck <chemin.json>
    --refresh <chemin.json>
    --dry-run
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from config.ai_generation import (
    AI_GENERATOR_VERSION,
    OPENAI_MODEL,
    OPENAI_REASONING_EFFORT,
    OPENAI_VERBOSITY,
    ALLOWED_RECHECK_STATUSES,
    MIN_LIVE_DESCRIPTION_CHARS,
    DEFAULT_BASE_CV_NAME,
    CANDIDATE_TRUTH,
    FORBIDDEN_CANDIDATE_CLAIM_TERMS,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
LOG_DIR = PROJECT_ROOT / "exports" / "logs"
APPLICATION_DIR = PROJECT_ROOT / "exports" / "applications"

# ---------------------------------------------------------------------------
# Structured Output schema
# ---------------------------------------------------------------------------

GENERATION_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "language": {"type": "string"},
        "cv": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "headline": {"type": "string"},
                "summary": {"type": "string"},
                "key_skills": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "experiences": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "additionalProperties": False,
                        "properties": {
                            "company": {"type": "string"},
                            "role": {"type": "string"},
                            "dates": {"type": "string"},
                            "bullets": {
                                "type": "array",
                                "items": {"type": "string"},
                            },
                        },
                        "required": ["company", "role", "dates", "bullets"],
                    },
                },
                "education": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "data_project": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "languages": {
                    "type": "array",
                    "items": {"type": "string"},
                },
            },
            "required": [
                "headline", "summary", "key_skills", "experiences",
                "education", "data_project", "languages",
            ],
        },
        "cover_letter": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "subject": {"type": "string"},
                "salutation": {"type": "string"},
                "paragraphs": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "closing": {"type": "string"},
            },
            "required": ["subject", "salutation", "paragraphs", "closing"],
        },
        "email": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "subject": {"type": "string"},
                "body": {"type": "string"},
            },
            "required": ["subject", "body"],
        },
        "alignment": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "top_matches": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "gaps_not_to_fake": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "truth_check": {
                    "type": "array",
                    "items": {"type": "string"},
                },
            },
            "required": ["top_matches", "gaps_not_to_fake", "truth_check"],
        },
        "claims_requiring_verification": {
            "type": "array",
            "items": {"type": "string"},
        },
    },
    "required": [
        "language", "cv", "cover_letter", "email", "alignment",
        "claims_requiring_verification",
    ],
}


# ---------------------------------------------------------------------------
# Utilitaires
# ---------------------------------------------------------------------------

def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize(value: Any) -> str:
    text = clean_text(value).lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text)


def latest_file(directory: Path, pattern: str) -> Path | None:
    paths = list(directory.glob(pattern))
    return max(paths, key=lambda p: p.stat().st_mtime) if paths else None


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def safe_filename(value: str, max_len: int = 90) -> str:
    text = unicodedata.normalize("NFKD", clean_text(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("._-")
    return (text[:max_len] or "document")


def extract_docx_text(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx absent. Installe-le avec : pip install -U python-docx"
        ) from exc

    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        txt = clean_text(p.text)
        if txt:
            parts.append(txt)
    for table in doc.tables:
        for row in table.rows:
            vals = [clean_text(cell.text) for cell in row.cells]
            if any(vals):
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def extract_contact(base_cv_text: str) -> dict:
    email_match = re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        base_cv_text,
    )
    phone_match = re.search(
        r"(?:\+32|0032|0)[\s./-]?\d(?:[\s./-]?\d{2,3}){2,4}",
        base_cv_text,
    )
    return {
        "name": CANDIDATE_TRUTH["identity"]["name"],
        "location": CANDIDATE_TRUTH["identity"]["location"],
        "email": email_match.group(0) if email_match else "",
        "phone": phone_match.group(0) if phone_match else "",
    }


def index_by_key(items: list[dict]) -> dict[str, dict]:
    return {
        clean_text(x.get("stable_item_key")): x
        for x in items
        if clean_text(x.get("stable_item_key"))
    }


def resolve_inputs(
    recheck_path: str | None = None,
    refresh_path: str | None = None,
) -> tuple[Path, Path, list[dict], list[dict]]:
    rp = Path(recheck_path) if recheck_path else latest_file(
        LOG_DIR, "application_recheck_v1_*.json"
    )
    fp = Path(refresh_path) if refresh_path else latest_file(
        LOG_DIR, "job_refresh_v1_*.json"
    )

    if not rp or not rp.exists():
        raise RuntimeError("Aucun application_recheck_v1_*.json trouvé.")
    if not fp or not fp.exists():
        raise RuntimeError("Aucun job_refresh_v1_*.json trouvé.")

    recheck = load_json(rp)
    refresh = load_json(fp)
    if not isinstance(recheck, list) or not isinstance(refresh, list):
        raise RuntimeError("JSON recheck/refresh invalide.")

    return rp, fp, recheck, refresh


def select_jobs(
    recheck: list[dict],
    refresh: list[dict],
    limit: int | None = None,
) -> list[dict]:
    refresh_map = index_by_key(refresh)
    selected = []

    for decision in sorted(
        recheck,
        key=lambda x: int(x.get("queue_rank") or 999999),
    ):
        if clean_text(decision.get("status")) not in ALLOWED_RECHECK_STATUSES:
            continue

        key = clean_text(decision.get("stable_item_key"))
        live = refresh_map.get(key)
        if not live:
            continue
        if clean_text(live.get("job_live_status")) != "LIVE_CONFIRMED":
            continue

        live_text = clean_text((live.get("live_detail") or {}).get("matching_text"))
        if len(live_text) < MIN_LIVE_DESCRIPTION_CHARS:
            continue

        merged = {
            "decision": decision,
            "refresh": live,
            "stable_item_key": key,
            "title": live.get("title") or decision.get("title"),
            "company": live.get("company") or decision.get("company"),
            "location": live.get("location") or decision.get("location"),
            "url": live.get("url") or decision.get("url"),
            "cv_track": decision.get("cv_track") or live.get("cv_track"),
            "live_description": live_text,
            "application_folder": live.get("application_folder"),
            "output_files": live.get("output_files") or {},
            "base_cv_path": live.get("base_cv_path"),
        }
        selected.append(merged)

    if limit is not None:
        selected = selected[: max(0, int(limit))]
    return selected


# ---------------------------------------------------------------------------
# Prompt
# ---------------------------------------------------------------------------

def build_prompt(job: dict, base_cv_text: str) -> str:
    return f"""
Tu es un recruteur senior belge et rédacteur ATS.

OBJECTIF
Produire une candidature personnalisée, crédible et factuelle pour l'offre
ci-dessous, à partir EXCLUSIVEMENT :
1) des vérités candidat autorisées ;
2) du CV source ;
3) de la description live de l'offre.

RÈGLE ABSOLUE : AUCUNE INVENTION.
- N'invente jamais un outil, une méthode, une durée, un diplôme, une langue,
  une responsabilité ou une expérience.
- Ne transforme jamais une exigence du poste en compétence du candidat.
- Si une compétence manque, place-la uniquement dans gaps_not_to_fake.
- Si tu hésites sur un fait, ajoute-le à claims_requiring_verification.
- claims_requiring_verification doit être [] si tout ce que tu écris est
  solidement supporté.
- Le diplôme Business Data Analysis est OBTENU en 2026.
- Le candidat n'a PAS de diplôme de Master.
- Il a validé un niveau Master 1 en sciences pharmaceutiques, sans obtenir le
  Master.
- Il a 3 ans d'expérience professionnelle Pharma QC.
- Il a 0 an d'expérience professionnelle Data/BI ; le Data vient de sa
  formation/projets.
- HPLC/UPLC est supporté par le stage Corden Pharma.
- Ne revendique pas GC, Empower, HACCP, Talend, Apache Hop, Tableau, Oracle,
  PL/SQL ou agrément médical sauf si le profil autorisé le prouve, ce qui
  n'est actuellement pas le cas.
- Ne gonfle pas les expériences de stage en expérience professionnelle longue.

STYLE
- Français professionnel naturel, belge, direct et crédible.
- ATS-friendly : titres simples, mots-clés réellement supportés.
- Pas de flatterie excessive.
- Pas de phrases génériques vides.
- CV concis, idéalement 2 pages une fois rendu.
- Lettre : 3 à 5 paragraphes courts.
- Email : bref.
- Réutilise le vocabulaire de l'offre lorsque le fait candidat correspondant
  existe réellement.

VÉRITÉS CANDIDAT AUTORISÉES
{json.dumps(CANDIDATE_TRUTH, ensure_ascii=False, indent=2)}

CV SOURCE LOCAL
----------------
{base_cv_text}

OFFRE LIVE
----------
Titre : {job['title']}
Entreprise : {job['company']}
Lieu : {job['location']}
URL : {job['url']}
Track : {job['cv_track']}

Description complète :
{job['live_description']}

CONSIGNES CV
- Conserver les entreprises, dates et nature réelle des expériences.
- Réordonner les compétences et reformuler les bullets pour l'offre.
- Ne pas supprimer les expériences majeures.
- Pour un poste Lab/QC, prioriser Pharma QC / laboratoire.
- Pour un poste Data, prioriser BDA / Python / SQL / Power BI / ETL, tout en
  indiquant honnêtement que l'expérience Data est académique/projet.
- Éducation : corriger impérativement le BDA en diplôme obtenu en 2026.

CONSIGNES LETTRE
- Expliquer POURQUOI le parcours est pertinent pour cette offre précise.
- Ne pas prétendre satisfaire une exigence non démontrée.
- Si le profil présente un écart non bloquant, le traiter avec honnêteté sans
  s'auto-saboter.

Avant de répondre, vérifie mentalement chaque affirmation candidate contre les
sources. Retourne uniquement le JSON conforme au schéma.
""".strip()


# ---------------------------------------------------------------------------
# API OpenAI
# ---------------------------------------------------------------------------

def call_openai(
    prompt: str,
    model: str = OPENAI_MODEL,
) -> tuple[dict, dict]:
    if not os.getenv("OPENAI_API_KEY"):
        raise RuntimeError(
            "OPENAI_API_KEY absent. Configure la variable d'environnement "
            "avant de lancer la génération."
        )

    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError(
            "SDK openai absent. Installe-le avec : pip install -U openai"
        ) from exc

    client = OpenAI()

    response = client.responses.create(
        model=model,
        instructions=(
            "Tu dois respecter strictement les faits utilisateur et le schéma JSON. "
            "Toute invention de compétence, diplôme ou expérience invalide la candidature."
        ),
        input=prompt,
        reasoning={"effort": OPENAI_REASONING_EFFORT},
        text={
            "verbosity": OPENAI_VERBOSITY,
            "format": {
                "type": "json_schema",
                "name": "job_application_documents",
                "strict": True,
                "schema": GENERATION_SCHEMA,
            },
        },
    )

    raw = response.output_text
    if not raw:
        raise RuntimeError("Réponse OpenAI vide.")

    payload = json.loads(raw)

    usage = {}
    try:
        if response.usage:
            usage = response.usage.model_dump()
    except Exception:
        try:
            usage = dict(response.usage or {})
        except Exception:
            usage = {}

    meta = {
        "model": model,
        "response_id": getattr(response, "id", None),
        "usage": usage,
    }
    return payload, meta


# ---------------------------------------------------------------------------
# Validation factuelle locale
# ---------------------------------------------------------------------------

EXPECTED_EXPERIENCE_DATES = {
    "gsk": "05/2022",
    "prothya": "06/2020",
    "corden pharma": "01/2017",
    "institut meurice": "04/2017",
}


def candidate_text_only(payload: dict) -> str:
    cv = payload.get("cv") or {}
    cover = payload.get("cover_letter") or {}
    email = payload.get("email") or {}

    parts = [
        cv.get("headline"),
        cv.get("summary"),
        " ".join(cv.get("key_skills") or []),
        " ".join(cv.get("education") or []),
        " ".join(cv.get("data_project") or []),
        " ".join(cv.get("languages") or []),
        cover.get("subject"),
        cover.get("salutation"),
        " ".join(cover.get("paragraphs") or []),
        cover.get("closing"),
        email.get("subject"),
        email.get("body"),
    ]
    for exp in cv.get("experiences") or []:
        parts.extend([
            exp.get("company"),
            exp.get("role"),
            exp.get("dates"),
            " ".join(exp.get("bullets") or []),
        ])
    return "\n".join(clean_text(x) for x in parts if clean_text(x))


def validate_generated_payload(payload: dict) -> list[str]:
    errors = []

    claims = payload.get("claims_requiring_verification") or []
    if claims:
        errors.append(
            "Le modèle a signalé des affirmations à vérifier : "
            + " | ".join(clean_text(x) for x in claims)
        )

    cv = payload.get("cv") or {}
    exps = cv.get("experiences") or []
    exp_text = normalize(" ".join(
        f"{e.get('company','')} {e.get('dates','')}" for e in exps
    ))

    for company, expected_start in EXPECTED_EXPERIENCE_DATES.items():
        if company not in exp_text:
            errors.append(f"Expérience attendue absente : {company}")
        if normalize(expected_start) not in exp_text:
            errors.append(
                f"Date de départ attendue non retrouvée pour {company}: "
                f"{expected_start}"
            )

    all_text = normalize(candidate_text_only(payload))

    # Le diplôme BDA doit être présenté comme obtenu.
    if "business data" in all_text and any(
        phrase in all_text for phrase in [
            "en cours", "en finalisation", "ongoing", "in progress"
        ]
    ):
        errors.append("Le BDA est encore présenté comme en cours/finalisation.")

    # Ne jamais revendiquer un Master obtenu.
    # Bloquer uniquement une revendication POSITIVE de Master.
    # Les formulations factuelles "sans diplôme de Master" / "niveau Master 1"
    # doivent rester autorisées.
    positive_master_patterns = [
        r"\btitulaire d['’]un master\b",
        r"\bmaster obtenu\b",
        r"\bdiplome de master obtenu\b",
        r"\bdiplomee? d['’]un master\b",
        r"\bmaster en sciences pharmaceutiques\b(?!.*sans diplome)",
    ]
    if any(re.search(pattern, all_text) for pattern in positive_master_patterns):
        errors.append("Le texte semble revendiquer un diplôme de Master.")

    # Termes unsupported particulièrement dangereux.
    for needle, label in FORBIDDEN_CANDIDATE_CLAIM_TERMS.items():
        if normalize(needle) in all_text:
            errors.append(
                f"Terme non prouvé présent dans les documents candidat : {label}"
            )

    # Interdire les affirmations évidentes de GC / HACCP comme maîtrise acquise.
    suspicious_patterns = [
        r"\b(maitrise|maitriser|experience|experimente|expertise)\b.{0,20}\bgc\b",
        r"\bgc\b.{0,20}\b(maitrise|experience|expertise)\b",
        r"\b(maitrise|experience|expertise|connaissance pratique)\b.{0,25}\bhaccp\b",
    ]
    for pattern in suspicious_patterns:
        if re.search(pattern, all_text):
            errors.append(
                "Le texte semble revendiquer GC/HACCP sans preuve suffisante."
            )
            break

    return errors


# ---------------------------------------------------------------------------
# DOCX local
# ---------------------------------------------------------------------------

def _configure_document(doc):
    from docx.shared import Cm, Pt

    sec = doc.sections[0]
    sec.top_margin = Cm(1.45)
    sec.bottom_margin = Cm(1.45)
    sec.left_margin = Cm(1.55)
    sec.right_margin = Cm(1.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)


def _add_section_title(doc, text):
    from docx.shared import Pt, RGBColor

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(31, 78, 121)


def _add_bullets(doc, bullets, size=9.0):
    from docx.shared import Cm, Pt

    for bullet in bullets:
        txt = clean_text(bullet)
        if not txt:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.18)
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run("• " + txt)
        r.font.size = Pt(size)


def build_cv_docx(payload: dict, contact: dict, path: Path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    cv = payload["cv"]
    doc = Document()
    _configure_document(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(contact["name"])
    r.bold = True
    r.font.size = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(clean_text(cv["headline"]).upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(31, 78, 121)

    contact_line = " | ".join(
        x for x in [
            contact.get("location"),
            contact.get("phone"),
            contact.get("email"),
        ] if clean_text(x)
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(contact_line).font.size = Pt(8.7)

    _add_section_title(doc, "Profil")
    doc.add_paragraph(clean_text(cv["summary"]))

    _add_section_title(doc, "Compétences clés")
    _add_bullets(doc, cv["key_skills"], size=8.9)

    _add_section_title(doc, "Expérience professionnelle")
    for exp in cv["experiences"]:
        p = doc.add_paragraph()
        r = p.add_run(
            f"{clean_text(exp['role'])} — {clean_text(exp['company'])}"
        )
        r.bold = True
        r.font.size = Pt(9.7)
        r2 = p.add_run(f"  |  {clean_text(exp['dates'])}")
        r2.italic = True
        r2.font.size = Pt(8.7)
        _add_bullets(doc, exp["bullets"], size=8.7)

    _add_section_title(doc, "Formation")
    for item in cv["education"]:
        doc.add_paragraph(clean_text(item))

    if cv.get("data_project"):
        _add_section_title(doc, "Projet / compétences Data")
        _add_bullets(doc, cv["data_project"], size=8.8)

    _add_section_title(doc, "Langues")
    doc.add_paragraph(" | ".join(clean_text(x) for x in cv["languages"]))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def build_cover_letter_docx(
    payload: dict,
    contact: dict,
    job: dict,
    path: Path,
):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    letter = payload["cover_letter"]
    doc = Document()
    _configure_document(doc)

    p = doc.add_paragraph()
    r = p.add_run(contact["name"])
    r.bold = True
    p.add_run(
        "\n" + "\n".join(
            x for x in [
                contact.get("location"),
                contact.get("phone"),
                contact.get("email"),
            ] if clean_text(x)
        )
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(clean_text(job["company"]))
    r.bold = True
    if clean_text(job["location"]):
        p.add_run("\n" + clean_text(job["location"]))

    p = doc.add_paragraph()
    p.add_run(datetime.now().strftime("Bruxelles, le %d/%m/%Y"))

    p = doc.add_paragraph()
    r = p.add_run(clean_text(letter["subject"]))
    r.bold = True

    doc.add_paragraph(clean_text(letter["salutation"]))

    for paragraph in letter["paragraphs"]:
        p = doc.add_paragraph(clean_text(paragraph))
        p.paragraph_format.space_after = Pt(7)

    doc.add_paragraph(clean_text(letter["closing"]))

    p = doc.add_paragraph()
    r = p.add_run(contact["name"])
    r.bold = True

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def output_paths(job: dict) -> tuple[Path, Path, Path, Path]:
    folder_value = clean_text(job.get("application_folder"))
    if folder_value:
        folder = PROJECT_ROOT / Path(folder_value)
    else:
        folder = APPLICATION_DIR / (
            safe_filename(job["stable_item_key"]) + "_"
            + safe_filename(job["company"]) + "_"
            + safe_filename(job["title"])
        )

    out_files = job.get("output_files") or {}
    cv_name = clean_text(out_files.get("cv_docx")) or "CV.docx"
    letter_name = (
        clean_text(out_files.get("cover_letter_docx"))
        or "Lettre_Motivation.docx"
    )
    email_name = clean_text(out_files.get("email_txt")) or "Email.txt"
    ai_json_name = safe_filename(job["stable_item_key"]) + "_AI_Generation.json"

    return (
        folder / cv_name,
        folder / letter_name,
        folder / email_name,
        folder / ai_json_name,
    )


def write_job_outputs(
    job: dict,
    payload: dict,
    meta: dict,
    contact: dict,
) -> dict:
    cv_path, letter_path, email_path, json_path = output_paths(job)

    build_cv_docx(payload, contact, cv_path)
    build_cover_letter_docx(payload, contact, job, letter_path)

    email = payload["email"]
    email_path.parent.mkdir(parents=True, exist_ok=True)
    email_path.write_text(
        f"Objet : {clean_text(email['subject'])}\n\n{clean_text(email['body'])}\n",
        encoding="utf-8",
    )

    generation_record = {
        "generator_version": AI_GENERATOR_VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "stable_item_key": job["stable_item_key"],
        "title": job["title"],
        "company": job["company"],
        "location": job["location"],
        "url": job["url"],
        "model": meta.get("model"),
        "response_id": meta.get("response_id"),
        "usage": meta.get("usage") or {},
        "payload": payload,
        "files": {
            "cv_docx": str(cv_path),
            "cover_letter_docx": str(letter_path),
            "email_txt": str(email_path),
            "generation_json": str(json_path),
        },
    }
    json_path.write_text(
        json.dumps(generation_record, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return generation_record


def export_run_log(records: list[dict], failures: list[dict]) -> tuple[Path, Path]:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    txt_path = LOG_DIR / f"ai_document_generator_v1_{stamp}.txt"
    json_path = LOG_DIR / f"ai_document_generator_v1_{stamp}.json"

    run = {
        "generator_version": AI_GENERATOR_VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "success_count": len(records),
        "failure_count": len(failures),
        "records": records,
        "failures": failures,
    }
    json_path.write_text(
        json.dumps(run, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    lines = [
        "AI DOCUMENT GENERATOR V1",
        "=" * 78,
        f"Succès : {len(records)}",
        f"Échecs : {len(failures)}",
        "",
    ]
    for i, rec in enumerate(records, 1):
        lines.extend([
            f"{i:2d}. OK | {rec.get('company')} | {rec.get('title')}",
            f"    Modèle : {rec.get('model')}",
            f"    CV     : {(rec.get('files') or {}).get('cv_docx')}",
            f"    Lettre : {(rec.get('files') or {}).get('cover_letter_docx')}",
            "",
        ])
    for fail in failures:
        lines.extend([
            f"❌ {fail.get('stable_item_key')} | {fail.get('title')}",
            f"   {fail.get('error')}",
            "",
        ])

    txt_path.write_text("\n".join(lines), encoding="utf-8")
    return txt_path, json_path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--model", default=OPENAI_MODEL)
    parser.add_argument("--recheck", dest="recheck_path", default=None)
    parser.add_argument("--refresh", dest="refresh_path", default=None)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    rp, fp, recheck, refresh = resolve_inputs(
        args.recheck_path,
        args.refresh_path,
    )
    jobs = select_jobs(recheck, refresh, args.limit)

    if not jobs:
        raise SystemExit("Aucune offre READY_DOCUMENTS exploitable.")

    # Le CV peut varier d'un item à l'autre, mais dans le projet actuel le
    # même CV maître est utilisé. On relit explicitement le chemin du premier.
    base_cv_raw = clean_text(jobs[0].get("base_cv_path"))
    base_cv_path = Path(base_cv_raw) if base_cv_raw else (
        PROJECT_ROOT / "assets" / "cv" / DEFAULT_BASE_CV_NAME
    )
    if not base_cv_path.exists():
        raise SystemExit(f"CV de base introuvable : {base_cv_path}")

    base_cv_text = extract_docx_text(base_cv_path)
    contact = extract_contact(base_cv_text)

    print("=" * 78)
    print("AI DOCUMENT GENERATOR V1")
    print("=" * 78)
    print(f"Recheck : {rp}")
    print(f"Refresh : {fp}")
    print(f"CV      : {base_cv_path}")
    print(f"Modèle  : {args.model}")
    print(f"Offres  : {len(jobs)}")
    print(f"Mode    : {'DRY-RUN' if args.dry_run else 'API + DOCX'}")
    print()

    if args.dry_run:
        for i, job in enumerate(jobs, 1):
            prompt = build_prompt(job, base_cv_text)
            print(
                f"{i:2d}. {job['company']} | {job['title']} | "
                f"prompt={len(prompt)} caractères"
            )
        return

    if not os.getenv("OPENAI_API_KEY"):
        raise SystemExit(
            "OPENAI_API_KEY absent. Configure-la dans PowerShell avant ce run."
        )

    records = []
    failures = []

    for i, job in enumerate(jobs, 1):
        print(f"[{i}/{len(jobs)}] {job['company']} | {job['title']}")
        try:
            prompt = build_prompt(job, base_cv_text)
            payload, meta = call_openai(prompt, model=args.model)

            validation_errors = validate_generated_payload(payload)
            if validation_errors:
                raise RuntimeError(
                    "FACT_GUARD: " + " || ".join(validation_errors)
                )

            record = write_job_outputs(job, payload, meta, contact)
            records.append(record)
            print("    ✅ Documents générés")
        except Exception as exc:
            failures.append({
                "stable_item_key": job["stable_item_key"],
                "title": job["title"],
                "company": job["company"],
                "error": f"{type(exc).__name__}: {exc}",
            })
            print(f"    ❌ {type(exc).__name__}: {exc}")

    txt_path, json_path = export_run_log(records, failures)

    print()
    print("=" * 78)
    print("BILAN")
    print("=" * 78)
    print(f"Succès : {len(records)}")
    print(f"Échecs : {len(failures)}")
    print(f"TXT    : {txt_path}")
    print(f"JSON   : {json_path}")


if __name__ == "__main__":
    main()
